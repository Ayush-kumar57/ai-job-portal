# parser.py - resume file (PDF / DOCX) ko ek lambi text string mein badalta hai
import pymupdf   # PyMuPDF (pehle iska naam fitz tha)
import docx      # python-docx


def extract_text(path):
    lower = path.lower()

    if lower.endswith('.pdf'):
        with pymupdf.open(path) as doc:
            return '\n'.join(page.get_text() for page in doc).strip()

    if lower.endswith('.docx'):
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        # Bahut resumes mein skills/education table ke andar hote hain
        for table in d.tables:
            for row in table.rows:
                parts.append(' '.join(cell.text for cell in row.cells))
        return '\n'.join(parts).strip()

    raise ValueError('Unsupported file type')
